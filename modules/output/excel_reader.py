"""Read a customer-curated Excel workbook back into a filtered inventory.

Reads the Dimensions, Facts, and Metrics tabs, filters to rows where
"Include in SV?" == "Yes", and returns an inventory dict compatible
with the existing YAML generator.

Requires: openpyxl
"""

from typing import Any

from ..common.logger import get_logger

log = get_logger("excel_reader")


def read_curated_workbook(
    workbook_path: str,
    include_review: bool = False,
) -> dict:
    """Read a curated Excel workbook and return a filtered inventory.

    Args:
        workbook_path: Path to the .xlsx file.
        include_review: If True, also include items marked "Review".

    Returns:
        Inventory dict with only the selected items.
    """
    import openpyxl

    log.info("Reading curated workbook: %s", workbook_path)
    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)

    accepted = {"yes"}
    if include_review:
        accepted.add("review")

    dims = _read_item_tab(wb, "Dimensions", accepted)
    facts = _read_item_tab(wb, "Facts", accepted)
    metrics = _read_item_tab(wb, "Metrics", accepted)
    rels = _read_relationships_tab(wb)
    tables = _read_tables_tab(wb)

    # Collect referenced table names from selected items
    referenced_tables = set()
    for item in dims + facts + metrics:
        if item.get("table"):
            referenced_tables.add(item["table"])

    # Filter tables to only those referenced
    if tables:
        tables = [t for t in tables if t.get("name") in referenced_tables]

    # Filter relationships to only those between referenced tables
    if rels:
        rels = [r for r in rels if r.get("left_table") in referenced_tables
                and r.get("right_table") in referenced_tables]

    inventory = {
        "source_type": "curated",
        "snowflake_target": {},
        "tables": tables,
        "relationships": rels,
        "dimensions": dims,
        "facts": facts,
        "metrics": metrics,
        "filters": [],
        "flagged": [],
        "complexity_summary": _tally_complexity(dims + facts + metrics),
        "errors": [],
    }

    total = len(dims) + len(facts) + len(metrics)
    log.info(
        "Curated inventory: %d tables, %d dims, %d facts, %d metrics (%d total items).",
        len(tables), len(dims), len(facts), len(metrics), total,
    )

    wb.close()
    return inventory


# ---------------------------------------------------------------------------
# Tab readers
# ---------------------------------------------------------------------------

def _read_item_tab(wb, tab_name: str, accepted: set[str]) -> list[dict]:
    """Read Dimensions/Facts/Metrics tab, filtering by Include column."""
    if tab_name not in wb.sheetnames:
        log.warning("Tab '%s' not found in workbook.", tab_name)
        return []

    ws = wb[tab_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]

    # Find column indices
    col_map = {}
    for i, h in enumerate(headers):
        if "table" in h:
            col_map["table"] = i
        elif "name" == h or h == "name":
            col_map["name"] = i
        elif "expression" in h or "expr" in h:
            col_map["expr"] = i
        elif "data" in h and "type" in h:
            col_map["data_type"] = i
        elif "description" in h or "desc" in h:
            col_map["description"] = i
        elif "complexity" in h:
            col_map["complexity"] = i
        elif "include" in h:
            col_map["include"] = i

    if "include" not in col_map:
        log.warning("No 'Include in SV?' column found in '%s' tab. Including all rows.", tab_name)

    items = []
    for row in rows[1:]:
        if not row or all(v is None for v in row):
            continue

        # Check include column
        if "include" in col_map:
            inc_val = str(row[col_map["include"]] or "").strip().lower()
            if inc_val not in accepted:
                continue

        item = {
            "name": str(row[col_map.get("name", 1)] or ""),
            "expr": str(row[col_map.get("expr", 2)] or ""),
            "data_type": str(row[col_map.get("data_type", 3)] or "VARCHAR"),
            "table": str(row[col_map.get("table", 0)] or ""),
            "description": str(row[col_map.get("description", 4)] or ""),
            "synonyms": [],
            "complexity": str(row[col_map.get("complexity", 5)] or "simple"),
        }
        items.append(item)

    return items


def _read_relationships_tab(wb) -> list[dict]:
    """Read the Relationships tab."""
    if "Relationships" not in wb.sheetnames:
        return []

    ws = wb["Relationships"]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []

    rels = []
    for row in rows[1:]:
        if not row or all(v is None for v in row):
            continue
        rels.append({
            "left_table": str(row[0] or ""),
            "left_column": str(row[1] or ""),
            "right_table": str(row[2] or ""),
            "right_column": str(row[3] or ""),
            "relationship_type": str(row[4] or "many:one"),
            "is_active": str(row[5] or "Yes").lower() == "yes",
        })
    return rels


def _read_tables_tab(wb) -> list[dict]:
    """Read the Tables tab."""
    if "Tables" not in wb.sheetnames:
        return []

    ws = wb["Tables"]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []

    tables = []
    for row in rows[1:]:
        if not row or all(v is None for v in row):
            continue
        tables.append({
            "name": str(row[0] or ""),
            "physical_table": str(row[1] or str(row[0] or "")),
        })
    return tables


def _tally_complexity(items: list[dict]) -> dict:
    """Tally complexity counts."""
    summary = {"simple": 0, "needs_translation": 0, "manual_required": 0}
    for item in items:
        c = item.get("complexity", "simple")
        if c in summary:
            summary[c] += 1
    return summary
